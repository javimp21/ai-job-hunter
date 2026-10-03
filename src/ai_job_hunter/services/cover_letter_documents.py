"""Render a cover-letter draft as a Word (.docx) or PDF file for human review.

Both layouts are identical: applicant name, one contact line built only from
the configured facts, a localized date, a blank line and the letter body.
Nothing here sends or uploads anything.
"""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import TYPE_CHECKING
from xml.sax.saxutils import escape

from ai_job_hunter.application_prep.configuration import CandidateApplicationFacts

if TYPE_CHECKING:
    from ai_job_hunter.services.cover_letters import CoverLetterDraft

_SPANISH_MONTHS = (
    "enero", "febrero", "marzo", "abril", "mayo", "junio",
    "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
)
_ENGLISH_MONTHS = (
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
)
_BODY_POINTS = 11
_NAME_POINTS = 16


def applicant_name(applicant: CandidateApplicationFacts) -> str:
    return " ".join(part for part in (applicant.first_name, applicant.last_name) if part)


def contact_line(applicant: CandidateApplicationFacts) -> str:
    profile = _short_url(applicant.linkedin_url) if applicant.linkedin_url else None
    parts = (applicant.email, applicant.phone, applicant.current_city, profile)
    return " · ".join(part for part in parts if part)


def _short_url(url: str) -> str:
    """Drop scheme, "www." and trailing slash so the contact line fits on one line."""

    return re.sub(r"^https?://(?:www\.)?", "", url.strip()).rstrip("/")


def format_date(value: date, language: str) -> str:
    if language == "es":
        return f"{value.day} de {_SPANISH_MONTHS[value.month - 1]} de {value.year}"
    return f"{value.day} {_ENGLISH_MONTHS[value.month - 1]} {value.year}"


def body_paragraphs(text: str) -> list[list[str]]:
    """Paragraphs split on blank lines; each is a list of lines (single newlines)."""

    paragraphs = [block for block in re.split(r"\n\s*\n", text.replace("\r\n", "\n").strip()) if block.strip()]
    return [[line.strip() for line in block.split("\n")] for block in paragraphs]


def render_docx(
    draft: CoverLetterDraft, applicant: CandidateApplicationFacts, path: Path, *, today: date | None = None
) -> Path:
    from docx import Document
    from docx.shared import Cm, Pt

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.5)
    section.top_margin = section.bottom_margin = Cm(2.5)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(_BODY_POINTS)
    normal.paragraph_format.space_after = Pt(0)

    name = applicant_name(applicant)
    if name:
        run = document.add_paragraph().add_run(name)
        run.bold = True
        run.font.size = Pt(_NAME_POINTS)
    contact = contact_line(applicant)
    if contact:
        document.add_paragraph(contact)
    document.add_paragraph(format_date(today or date.today(), draft.language))
    document.add_paragraph()
    for lines in body_paragraphs(draft.text):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(_BODY_POINTS)
        for index, line in enumerate(lines):
            run = paragraph.add_run(line)
            if index < len(lines) - 1:
                run.add_break()

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def render_pdf(
    draft: CoverLetterDraft, applicant: CandidateApplicationFacts, path: Path, *, today: date | None = None
) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    body = ParagraphStyle("Body", fontName="Helvetica", fontSize=_BODY_POINTS, leading=_BODY_POINTS * 1.4)
    heading = ParagraphStyle(
        "Name", parent=body, fontName="Helvetica-Bold", fontSize=_NAME_POINTS, leading=_NAME_POINTS * 1.3
    )
    paragraph_style = ParagraphStyle("Paragraph", parent=body, spaceAfter=_BODY_POINTS)

    story: list = []
    name = applicant_name(applicant)
    if name:
        story.append(Paragraph(_markup(name), heading))
    contact = contact_line(applicant)
    if contact:
        story.append(Paragraph(_markup(contact), body))
    story.append(Paragraph(_markup(format_date(today or date.today(), draft.language)), body))
    story.append(Spacer(1, _BODY_POINTS * 1.4))
    for lines in body_paragraphs(draft.text):
        story.append(Paragraph("<br/>".join(_markup(line) for line in lines), paragraph_style))

    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=2.5 * cm,
        rightMargin=2.5 * cm,
        topMargin=2.5 * cm,
        bottomMargin=2.5 * cm,
        title=f"{draft.company} - {draft.title}",
        author=name or "",
    ).build(story)
    return path


def _markup(value: str) -> str:
    """Escape for Paragraph markup; built-in fonts cover cp1252 (accents, ñ, €, curly quotes, dashes)."""

    return escape(value.encode("cp1252", errors="replace").decode("cp1252"))
