"""Render a CV written in a small Markdown subset into ATS-friendly Word and PDF.

Supported lines: ``# Name``, ``## SECTION``, ``- bullet`` and plain text, with
``**bold**`` and ``*italic*`` inline. Output is single column, real text and
standard bullets (Word's "List Bullet" style, "•" in the PDF) so applicant
tracking systems extract it cleanly.
"""

from __future__ import annotations

import re
from pathlib import Path

CvBlock = tuple[str, str]  # (kind, text): kind is name, section, bullet or text
_NAVY = "1F3864"
_RULE = "8EA9C1"
_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def parse_cv(markdown: str) -> list[CvBlock]:
    blocks: list[CvBlock] = []
    for line in markdown.splitlines():
        line = line.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            blocks.append(("name", line[2:].strip()))
        elif line.startswith("## "):
            blocks.append(("section", line[3:].strip()))
        elif line.startswith("- "):
            blocks.append(("bullet", line[2:].strip()))
        else:
            blocks.append(("text", line.strip()))
    return blocks


def _runs(text: str) -> list[tuple[str, bool, bool]]:
    runs = []
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            runs.append((part[2:-2], True, False))
        elif part.startswith("*"):
            runs.append((part[1:-1], False, True))
        else:
            runs.append((part, False, False))
    return runs


def render_cv_docx(blocks: list[CvBlock], path: Path) -> Path:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    section.left_margin = section.right_margin = Cm(1.8)
    section.top_margin = section.bottom_margin = Cm(1.5)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)
    for kind, text in blocks:
        if kind == "name":
            run = document.add_paragraph().add_run(text)
            run.bold, run.font.size, run.font.color.rgb = True, Pt(20), RGBColor.from_string(_NAVY)
            continue
        if kind == "section":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(8)
            run = paragraph.add_run(text)
            run.bold, run.font.size, run.font.color.rgb = True, Pt(11), RGBColor.from_string(_NAVY)
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", _RULE)):
                bottom.set(qn(key), value)
            border.append(bottom)
            paragraph._p.get_or_add_pPr().append(border)
            continue
        paragraph = document.add_paragraph(style="List Bullet" if kind == "bullet" else None)
        for value, bold, italic in _runs(text):
            run = paragraph.add_run(value)
            run.bold, run.italic = bold, italic
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def render_cv_pdf(blocks: list[CvBlock], path: Path, *, title: str = "CV") -> Path:
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import HRFlowable, ListFlowable, ListItem, Paragraph, SimpleDocTemplate

    text_style = ParagraphStyle("cv-text", fontName="Helvetica", fontSize=9.8, leading=12.6, spaceAfter=2)
    name_style = ParagraphStyle("cv-name", fontName="Helvetica-Bold", fontSize=19, leading=23, textColor=HexColor(f"#{_NAVY}"))
    section_style = ParagraphStyle(
        "cv-section", fontName="Helvetica-Bold", fontSize=10.5, leading=13, textColor=HexColor(f"#{_NAVY}"), spaceBefore=8
    )
    story: list = []
    pending: list[str] = []

    def html(value: str) -> str:
        value = value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        value = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", value)
        return re.sub(r"\*([^*]+)\*", r"<i>\1</i>", value)

    def flush() -> None:
        if pending:
            story.append(
                ListFlowable(
                    [ListItem(Paragraph(item, text_style), leftIndent=12) for item in pending],
                    bulletType="bullet", start="•", leftIndent=12, bulletFontSize=9,
                )
            )
            pending.clear()

    for kind, text in blocks:
        if kind == "bullet":
            pending.append(html(text))
            continue
        flush()
        if kind == "name":
            story.append(Paragraph(html(text), name_style))
        elif kind == "section":
            story.append(Paragraph(html(text), section_style))
            story.append(HRFlowable(width="100%", thickness=0.6, color=HexColor(f"#{_RULE}"), spaceBefore=1, spaceAfter=3))
        else:
            story.append(Paragraph(html(text), text_style))
    flush()
    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
        title=title,
    ).build(story)
    return path
