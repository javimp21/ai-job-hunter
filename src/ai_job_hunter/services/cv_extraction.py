"""Reads a CV (PDF, Word or pasted text) into a draft profile with a small, cheap Claude model.

The person confirms or corrects the draft in the sign-up conversation, so this only has to be good enough to save typing.
The CV is read once and never stored. The tool schema has no field for name, e-mail, phone or address on purpose.
"""

from __future__ import annotations

import base64
import io
from typing import Any

from ai_job_hunter.services.cover_letters import CoverLetterError, _anthropic_client
from ai_job_hunter.services.onboarding import ProfileExtractor

MODEL = "claude-haiku-5-5"
TOOL_NAME = "cv_profile"
MAX_TEXT_CHARS = 60_000
MAX_ITEMS = 25
MAX_ITEM_CHARS = 60

_SYSTEM = """You read a CV and fill in a short professional profile by calling the tool once.
- Use only what the CV says; leave a field empty when it does not say. Never guess.
- Do NOT output the person's name, e-mail, phone number or street address, and do not copy them anywhere.
- years_of_experience: total years of professional work experience as a number (internships count as half).
- primary_skills: the 3 to 8 core professional skills; technologies: tools, languages, software and standards named.
- languages: spoken languages, as English names ("Spanish", "English", "French").
- current_city and current_country: where the person lives now, if stated (city name only, no street).
If a previous draft and a correction from the person are given, apply the correction to the draft and keep the rest."""

_STRING_LISTS = ("primary_skills", "secondary_skills", "technologies", "languages")
_STRINGS = ("current_role", "education", "current_country", "current_city")


def _schema() -> dict[str, Any]:
    text = {"type": "string"}
    listing = {"type": "array", "items": {"type": "string"}}
    return {
        "type": "object",
        "properties": {
            "current_role": text, "years_of_experience": {"type": "number"}, "primary_skills": listing,
            "secondary_skills": listing, "technologies": listing, "languages": listing, "education": text,
            "current_country": text, "current_city": text,
        },
        "required": [],
    }


def docx_text(content: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
    return "\n".join(part for part in parts if part.strip())[:MAX_TEXT_CHARS]


def build_request(
    text: str | None, file: tuple[str, bytes] | None, previous: dict[str, Any] | None, correction: str | None
) -> dict[str, Any]:
    content: list[dict[str, Any]] = []
    if file is not None:
        name, data = file
        if name.casefold().endswith(".pdf"):
            content.append({"type": "document", "source": {
                "type": "base64", "media_type": "application/pdf", "data": base64.b64encode(data).decode("ascii")}})
        else:
            text = docx_text(data)
    if text:
        content.append({"type": "text", "text": f"<cv>\n{text[:MAX_TEXT_CHARS]}\n</cv>"})
    if previous and correction:
        content.append({"type": "text", "text": f"Previous draft: {previous}\nCorrection from the person: {correction}"})
    content.append({"type": "text", "text": "Fill in the profile."})
    return {
        "model": MODEL,
        "max_tokens": 1200,
        "system": _SYSTEM,
        "messages": [{"role": "user", "content": content}],
        "tools": [{"name": TOOL_NAME, "description": "The profile read from the CV.", "input_schema": _schema()}],
        "tool_choice": {"type": "tool", "name": TOOL_NAME},
    }


def normalize(raw: dict[str, Any]) -> dict[str, Any]:
    """Keep only expected fields, trimmed and capped; a bad value is dropped rather than trusted."""

    draft: dict[str, Any] = {}
    for key in _STRINGS:
        value = raw.get(key)
        if isinstance(value, str) and value.strip():
            draft[key] = value.strip()[:MAX_ITEM_CHARS * 3]
    for key in _STRING_LISTS:
        value = raw.get(key)
        if isinstance(value, list):
            items = [item.strip()[:MAX_ITEM_CHARS] for item in value if isinstance(item, str) and item.strip()]
            if items:
                draft[key] = list(dict.fromkeys(items))[:MAX_ITEMS]
    years = raw.get("years_of_experience")
    if isinstance(years, (int, float)) and not isinstance(years, bool) and 0 <= years <= 60:
        draft["years_of_experience"] = round(float(years), 1)
    return draft


def make_extractor(client: Any | None = None) -> ProfileExtractor:
    def extract(text, file, previous, correction) -> dict[str, Any]:
        active = client or _anthropic_client()
        response = active.messages.create(**build_request(text, file, previous, correction))
        block = next((item for item in response.content if getattr(item, "type", "") == "tool_use"), None)
        if block is None:
            raise CoverLetterError("The CV reader returned no profile.")
        draft = normalize(block.input)
        if not draft:
            raise CoverLetterError("The CV reader found nothing usable.")
        return draft

    return extract
