import base64
from datetime import UTC, date, datetime
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4

import pytest

from ai_job_hunter.application_prep.configuration import CandidateApplicationFacts
from ai_job_hunter.application_prep.models import CandidateDocument
from ai_job_hunter.candidates import CandidateConfig
from ai_job_hunter.models import Company, Job, JobSource
from ai_job_hunter.services.cover_letter_documents import render_docx, render_pdf
from ai_job_hunter.services.cover_letters import (
    COVER_LETTER_MODEL,
    CoverLetterDraft,
    CoverLetterError,
    build_cover_letter_request,
    candidate_facts_for_letter,
    generate_cover_letter,
)


def _candidate() -> CandidateConfig:
    return CandidateConfig.model_validate(
        {
            "profile": {
                "current_role": "Backend Engineer",
                "years_of_experience": 1,
                "current_salary": 99999,
                "salary_currency": "EUR",
                "primary_skills": ["Java", "Spring Boot"],
                "languages": ["Spanish", "English"],
            },
            "preferences": {
                "minimum_salary": 26000,
                "salary_currency": "EUR",
                "salary_period": "YEAR",
                "willing_to_learn_technologies": ["Kafka"],
            },
        }
    )


def _application() -> CandidateApplicationFacts:
    return CandidateApplicationFacts(
        first_name="Alex",
        last_name="Example",
        email="alex@example.test",
        phone="+34 600 000 000",
        current_company="Fictional Bank",
        linkedin_url="https://linkedin.example.test/alex",
    )


def _job(session) -> Job:
    company = Company(name="Acme", website_url="https://acme.example.test")
    job = Job(company=company, title="Backend Engineer")
    session.add_all([company, job])
    session.flush()
    session.add(
        JobSource(
            job=job,
            provider="greenhouse",
            external_id="role-1",
            original_url="https://boards.greenhouse.io/acme/jobs/1",
            canonical_url="https://boards.greenhouse.io/acme/jobs/1",
            source_title="Backend Engineer",
            source_description="Build Java services with Kafka. 2+ years preferred.",
            source_location="Remote Spain",
            remote_policy="REMOTE",
        )
    )
    session.commit()
    return job


class FakeMessages:
    def __init__(self, response=None, error: Exception | None = None):
        self.requests: list[dict] = []
        self.response = response
        self.error = error

    def create(self, **kwargs):
        self.requests.append(kwargs)
        if self.error is not None:
            raise self.error
        return self.response


def _client(text: str = "Hi,\n\nI build Java services.\n\nThanks,\nAlex", stop_reason: str = "end_turn", error=None):
    response = SimpleNamespace(
        content=[SimpleNamespace(type="thinking", thinking=""), SimpleNamespace(type="text", text=text)],
        stop_reason=stop_reason,
        model=COVER_LETTER_MODEL,
        usage=SimpleNamespace(input_tokens=1200, output_tokens=300),
    )
    messages = FakeMessages(response, error)
    return SimpleNamespace(beta=SimpleNamespace(messages=messages)), messages


def _style(tmp_path: Path) -> Path:
    path = tmp_path / "WRITING.md"
    path.write_text("# Style\n- Plain and honest.\n", encoding="utf-8")
    return path


def test_letter_facts_exclude_salary_and_private_contact_details():
    facts = candidate_facts_for_letter(_candidate(), _application())

    rendered = repr(facts)
    assert facts["first_name"] == "Alex"
    assert facts["current_company"] == "Fictional Bank"
    assert facts["technologies_wants_to_learn"] == ["Kafka"]
    for secret in ("99999", "26000", "alex@example.test", "600 000", "linkedin", "Example"):
        assert secret not in rendered


def test_request_uses_default_model_style_guide_cache_fallback_and_cv():
    request = build_cover_letter_request(
        posting={"company": "Acme", "title": "Backend Engineer", "description": "Build things."},
        candidate_facts={"first_name": "Alex"},
        style_guide="Never say 'passionate'.",
        cv_pdf=b"%PDF-fake",
    )

    assert request["model"] == "claude-opus-5-5"
    assert request["fallbacks"] == "default"
    assert request["betas"] == ["server-side-fallback-2026-07-01"]
    system = request["system"][0]
    assert system["text"].endswith("Never say 'passionate'.")
    assert system["cache_control"] == {"type": "ephemeral"}
    document, text = request["messages"][0]["content"]
    assert document["source"]["media_type"] == "application/pdf"
    assert base64.b64decode(document["source"]["data"]) == b"%PDF-fake"
    assert "<job_posting>" in text["text"] and "Build things." in text["text"]


def test_generate_saves_local_draft_and_returns_text(db_session, tmp_path):
    job = _job(db_session)
    client, messages = _client()
    cv = tmp_path / "cv.pdf"
    cv.write_bytes(b"%PDF-cv")

    draft = generate_cover_letter(
        db_session,
        _candidate(),
        job.id,
        application_facts=_application(),
        documents=(CandidateDocument(type="CV", id="cv", local_path=str(cv)),),
        style_guide_path=_style(tmp_path),
        output_dir=tmp_path / "letters",
        client=client,
        now=datetime(2026, 10, 3, 12, 0, tzinfo=UTC),
    )

    assert draft.text.startswith("Hi,")
    assert draft.company == "Acme" and draft.title == "Backend Engineer"
    assert (draft.input_tokens, draft.output_tokens) == (1200, 300)
    assert draft.path.name == "20261003-120000-acme-backend-engineer-en.md"
    assert draft.language == "en"
    saved = draft.path.read_text(encoding="utf-8")
    assert "DRAFT" in saved and "I build Java services." in saved
    assert "https://boards.greenhouse.io/acme/jobs/1" in saved
    request = messages.requests[0]
    assert request["messages"][0]["content"][0]["type"] == "document"
    assert "Kafka" in request["messages"][0]["content"][1]["text"]


def test_refusal_unknown_job_missing_style_and_provider_errors_are_safe(db_session, tmp_path):
    job = _job(db_session)
    style = _style(tmp_path)
    common = dict(application_facts=_application(), style_guide_path=style, output_dir=tmp_path / "out")

    refused, _ = _client(stop_reason="refusal")
    with pytest.raises(CoverLetterError, match="declined"):
        generate_cover_letter(db_session, _candidate(), job.id, client=refused, **common)

    with pytest.raises(CoverLetterError, match="Unknown job id"):
        generate_cover_letter(db_session, _candidate(), uuid4(), client=_client()[0], **common)

    with pytest.raises(CoverLetterError, match="style guide"):
        generate_cover_letter(
            db_session, _candidate(), job.id, client=_client()[0],
            application_facts=_application(), style_guide_path=tmp_path / "missing.md",
        )

    failing, _ = _client(error=RuntimeError("https://api.example.test?key=secret-value"))
    with pytest.raises(CoverLetterError) as error:
        generate_cover_letter(db_session, _candidate(), job.id, client=failing, **common)
    assert "secret-value" not in str(error.value)
    assert "RuntimeError" in str(error.value)
    assert not (tmp_path / "out").exists()


_POSTING = {"company": "Acme", "title": "Backend Engineer", "description": "Build things."}


def _request_text(language: str) -> str:
    request = build_cover_letter_request(
        posting=_POSTING, candidate_facts={"first_name": "Alex"}, style_guide="s", cv_pdf=None, language=language
    )
    return request["messages"][0]["content"][-1]["text"]


def test_language_instruction_per_option_and_invalid_value_rejected():
    assert "Language override" not in _request_text("auto")
    spanish = _request_text("es")
    assert "in Spanish" in spanish and "peninsular" in spanish and "tuteo" in spanish
    assert "in English" in _request_text("en")
    with pytest.raises(ValueError):
        build_cover_letter_request(
            posting=_POSTING, candidate_facts={}, style_guide="s", cv_pdf=None, language="fr"
        )


def test_generate_rejects_invalid_language_before_any_request(db_session, tmp_path):
    client, messages = _client()
    with pytest.raises(ValueError):
        generate_cover_letter(
            db_session, _candidate(), uuid4(), application_facts=_application(), client=client, language="fr"
        )
    assert messages.requests == []


def test_filename_and_header_include_language_and_forced_language_is_sent(db_session, tmp_path):
    job = _job(db_session)
    client, messages = _client(text="Hola,\n\nConstruyo servicios.\n\nUn saludo,\nAlex")

    draft = generate_cover_letter(
        db_session, _candidate(), job.id, application_facts=_application(),
        style_guide_path=_style(tmp_path), output_dir=tmp_path / "out", client=client,
        now=datetime(2026, 10, 3, 12, 0, tzinfo=UTC), language="es", render_documents=False,
    )

    assert draft.language == "es" and draft.path.name.endswith("-acme-backend-engineer-es.md")
    assert "- Language: es" in draft.path.read_text(encoding="utf-8")
    assert "in Spanish" in messages.requests[0]["messages"][0]["content"][-1]["text"]
    assert draft.docx_path is None and draft.pdf_path is None and draft.render_error is None


def test_auto_language_is_detected_from_the_letter(db_session, tmp_path):
    job = _job(db_session)
    client, _ = _client(text="Hola,\n\nMe interesa el puesto por la experiencia del equipo.\n\nUn saludo")
    draft = generate_cover_letter(
        db_session, _candidate(), job.id, application_facts=_application(),
        style_guide_path=_style(tmp_path), output_dir=tmp_path / "out", client=client, render_documents=False,
    )
    assert draft.language == "es"


def _draft(tmp_path: Path, text: str, language: str = "en"):
    return CoverLetterDraft(
        job_id=uuid4(), company="Acme", title="Backend Engineer", text=text,
        path=tmp_path / "x.md", model="m", input_tokens=None, output_tokens=None, language=language,
    )


def test_docx_has_name_contact_localized_date_and_body(tmp_path):
    from docx import Document

    text = "Hola,\n\nPrimer párrafo.\nSegunda línea.\n\nUn saludo,\nAlex"
    path = render_docx(_draft(tmp_path, text, "es"), _application(), tmp_path / "l.docx", today=date(2026, 10, 3))

    paragraphs = [p.text for p in Document(str(path)).paragraphs]
    assert paragraphs[0] == "Alex Example"
    assert paragraphs[1] == "alex@example.test · +34 600 000 000 · linkedin.example.test/alex"
    assert paragraphs[2] == "3 de octubre de 2026"
    assert paragraphs[3] == ""
    assert paragraphs[4:] == ["Hola,", "Primer párrafo.\nSegunda línea.", "Un saludo,\nAlex"]
    assert Document(str(path)).paragraphs[0].runs[0].bold is True

    english = render_docx(_draft(tmp_path, "Hi", "en"), _application(), tmp_path / "e.docx", today=date(2026, 10, 3))
    assert Document(str(english)).paragraphs[2].text == "3 October 2026"


def test_contact_line_uses_only_present_fields():
    from ai_job_hunter.services.cover_letter_documents import contact_line

    assert contact_line(CandidateApplicationFacts(first_name="A", current_city="Madrid")) == "Madrid"
    assert contact_line(CandidateApplicationFacts()) == ""


def test_documents_render_special_characters(tmp_path):
    text = "Señor Núñez: cobro 30.000 € — “bien” y ’ok’ <b>&amp;</b> ✓\n\nÁ É Í Ó Ú ñ"
    draft = _draft(tmp_path, text, "es")
    docx = render_docx(draft, _application(), tmp_path / "s.docx")
    pdf = render_pdf(draft, _application(), tmp_path / "s.pdf", today=date(2026, 10, 3))

    assert docx.stat().st_size > 0
    assert pdf.read_bytes().startswith(b"%PDF")


def test_generate_writes_docx_and_pdf_next_to_draft(db_session, tmp_path):
    job = _job(db_session)
    client, _ = _client()
    draft = generate_cover_letter(
        db_session, _candidate(), job.id, application_facts=_application(),
        style_guide_path=_style(tmp_path), output_dir=tmp_path / "out", client=client,
        now=datetime(2026, 10, 3, 12, 0, tzinfo=UTC),
    )
    assert draft.docx_path == draft.path.with_suffix(".docx") and draft.docx_path.is_file()
    assert draft.pdf_path == draft.path.with_suffix(".pdf")
    assert draft.pdf_path.read_bytes().startswith(b"%PDF")
    assert draft.render_error is None


def test_render_failure_keeps_text_draft(db_session, tmp_path, monkeypatch):
    def boom(*_args, **_kwargs):
        raise RuntimeError("secret detail")

    monkeypatch.setattr("ai_job_hunter.services.cover_letters.render_docx", boom)
    job = _job(db_session)
    client, _ = _client()
    draft = generate_cover_letter(
        db_session, _candidate(), job.id, application_facts=_application(),
        style_guide_path=_style(tmp_path), output_dir=tmp_path / "out", client=client,
    )
    assert draft.path.is_file() and draft.text.startswith("Hi,")
    assert draft.docx_path is None and draft.pdf_path is not None
    assert draft.render_error and "RuntimeError" in draft.render_error and "secret" not in draft.render_error
