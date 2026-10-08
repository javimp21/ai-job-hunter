import io
from types import SimpleNamespace

import pytest
from docx import Document
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from ai_job_hunter.db.base import Base
from ai_job_hunter.db.user_context import current_user_id
from ai_job_hunter.models import User
from ai_job_hunter.services.cv_extraction import build_request, make_extractor, normalize
from ai_job_hunter.services.telegram_bot import TelegramBotError, handle_update
from ai_job_hunter.services.telegram_signup import SignupHandlers
from ai_job_hunter.services.users import ensure_owner

OWNER_CHAT = "1"


class PersonBot:
    """A bot that records what it sends and serves a fixed CV file."""

    def __init__(self, sent, files, chat="?"):
        self.sent, self.files, self.chat = sent, files, chat

    def with_chat(self, chat_id):
        return PersonBot(self.sent, self.files, str(chat_id))

    def send_text(self, text, reply_markup=None, reply_to=None):
        self.sent.append((self.chat, text, reply_markup))
        return 500 + len(self.sent)

    def answer_callback_query(self, callback_id, text):
        pass

    def username(self):
        return "job_hunter_test_bot"

    def download_file(self, file_id, *, max_bytes):
        if file_id not in self.files:
            raise TelegramBotError(kind="FileUnavailableOrTooLarge")
        return self.files[file_id]


@pytest.fixture
def world():
    engine = create_engine("sqlite+pysqlite:///:memory:", poolclass=StaticPool, connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    with factory() as session:
        ensure_owner(session, telegram_chat_id=OWNER_CHAT)
        session.commit()

    def extract(text, file, previous, correction):
        return {"current_role": "Contable", "years_of_experience": 2, "languages": ["Spanish"]}

    sent, files, votes, notes = [], {"cv1": b"%PDF-1.4 fake"}, [], []

    def record_feedback(job_id, state, reason):
        votes.append((current_user_id(), job_id, state))

    def record_note(prompt_job_id, replied_message_id, text):
        notes.append((current_user_id(), replied_message_id, text))
        return True

    handlers = SignupHandlers(factory, extract, OWNER_CHAT, record_feedback, record_note)
    bot = PersonBot(sent, files, chat=OWNER_CHAT)

    def send(chat, text=None, *, data=None, document=None, kind="private", reply_to=None):
        message = {"message_id": 1, "chat": {"id": int(chat), "type": kind}, **({"reply_to_message": {"message_id": reply_to}} if reply_to else {})}
        if data is not None:
            update = {"update_id": 1, "callback_query": {"id": "c", "data": data, "message": message}}
        else:
            update = {"update_id": 1, "message": {**message, **({"text": text} if text else {}), **({"document": document} if document else {})}}
        return handle_update(update, chat_id=OWNER_CHAT, bot=bot, generate=lambda *a: None, generated={}, signup=handlers)

    return SimpleNamespace(factory=factory, sent=sent, send=send, votes=votes, notes=notes, handlers=handlers)


def last(world, chat):
    return [item for item in world.sent if item[0] == str(chat)][-1]


def test_the_owner_invites_and_a_new_person_signs_up_with_the_code(world) -> None:
    assert world.send(OWNER_CHAT, "/invite") == "invited"
    code_message = last(world, OWNER_CHAT)[1]
    code = code_message.split("Invitación (un solo uso, caduca en 7 días): ")[1][:8]
    assert f"/start {code}" in code_message and f"https://t.me/job_hunter_test_bot?start={code}" in code_message

    assert world.send("500", "hola") == "signup_refused" and "invitación" in last(world, 500)[1]
    assert world.send("500", "/start WRONG123") == "signup_refused" and "no existe" in last(world, 500)[1]
    assert world.send("500", f"/start {code}") == "signup_started"
    assert "¿Aceptas?" in last(world, 500)[1] and last(world, 500)[2]["inline_keyboard"][0][0]["callback_data"] == "ob:consent:yes"
    assert world.send("501", f"/start {code}") == "signup_refused" and "ya se usó" in last(world, 501)[1]

    assert world.send("500", data="ob:consent:yes") == "signup_step"
    assert world.send("500", document={"file_id": "cv1", "file_name": "cv.pdf", "mime_type": "application/pdf", "file_size": 2000}) == "signup_step"
    assert "Contable" in last(world, 500)[1]
    assert [item[1] for item in world.sent if item[0] == "500"].count("Leyendo tu CV…") == 1


def test_files_that_are_not_a_cv_or_are_too_big_are_refused_before_downloading(world) -> None:
    world.send(OWNER_CHAT, "/invite")
    code = last(world, OWNER_CHAT)[1].split(": ")[1][:8]
    world.send("500", f"/start {code}")
    world.send("500", data="ob:consent:yes")

    world.send("500", document={"file_id": "cv1", "file_name": "foto.png", "mime_type": "image/png", "file_size": 10})
    assert "PDF o Word" in last(world, 500)[1]
    world.send("500", document={"file_id": "cv1", "file_name": "cv.pdf", "mime_type": "application/pdf", "file_size": 9_000_000})
    assert "menos de 5 MB" in last(world, 500)[1]
    world.send("500", document={"file_id": "missing", "file_name": "cv.pdf", "mime_type": "application/pdf", "file_size": 10})
    assert "No he podido descargar" in last(world, 500)[1]


def test_only_private_chats_are_served_and_a_person_cannot_invite(world) -> None:
    assert world.send("-100", "/start ABCDEFGH", kind="group") == "ignored"
    world.send(OWNER_CHAT, "/invite")
    code = last(world, OWNER_CHAT)[1].split(": ")[1][:8]
    world.send("500", f"/start {code}")
    world.send("500", data="ob:consent:yes")
    before = len(world.sent)
    assert world.send("500", "/invite") == "signup_command"
    assert "No conozco ese comando" in last(world, 500)[1] or "/my_data" in last(world, 500)[1]
    assert not any("🎟" in item[1] for item in world.sent[before:])


def test_a_person_can_see_pause_and_erase_their_data_by_button(world) -> None:
    world.send(OWNER_CHAT, "/invite")
    code = last(world, OWNER_CHAT)[1].split(": ")[1][:8]
    world.send("500", f"/start {code}")
    world.send("500", data="ob:consent:yes")
    assert world.send("500", "/my_data") == "signup_command" and "Estado: ONBOARDING" in last(world, 500)[1]
    assert world.send("500", "/erase") == "signup_command" and last(world, 500)[2]
    assert world.send("500", data="ob:erase:yes") == "signup_erased"
    with world.factory() as session:
        person = session.scalar(select(User).where(User.is_owner.is_(False)))
        assert person.status == "DELETED" and person.telegram_chat_id is None
    assert world.send("500", "hola") == "signup_refused"  # the chat is free again and needs a new invitation


def _docx(text: str) -> bytes:
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


class FakeClaude:
    def __init__(self, answer):
        self.requests, self._answer = [], answer
        self.messages = SimpleNamespace(create=self._create)

    def _create(self, **request):
        self.requests.append(request)
        return SimpleNamespace(content=[SimpleNamespace(type="tool_use", input=self._answer)])


def test_the_cv_reader_sends_pdfs_natively_reads_word_files_and_never_asks_for_personal_contact_data() -> None:
    pdf = build_request(None, ("cv.pdf", b"%PDF"), None, None)
    assert pdf["model"] == "claude-haiku-5-5" and pdf["messages"][0]["content"][0]["type"] == "document"
    word = build_request(None, ("cv.docx", _docx("Contable con 2 años en Madrid")), None, None)
    assert "Contable con 2 años" in word["messages"][0]["content"][0]["text"]
    fields = set(pdf["tools"][0]["input_schema"]["properties"])
    assert not fields & {"name", "email", "phone", "address"} and "never" in pdf["system"].lower() or "Do NOT" in pdf["system"]
    corrected = build_request("cv", None, {"years_of_experience": 2}, "son 3")
    assert "Correction from the person: son 3" in corrected["messages"][0]["content"][1]["text"]


def test_the_cv_reader_cleans_what_the_model_returns_and_refuses_empty_results() -> None:
    draft = normalize({
        "current_role": " Contable ", "years_of_experience": 2.04, "languages": ["Spanish", "Spanish", "", 7],
        "technologies": ["Excel"] * 3 + ["x" * 200], "email": "a@b.c", "years_extra": 9,
    })
    assert draft["current_role"] == "Contable" and draft["years_of_experience"] == 2.0
    assert draft["languages"] == ["Spanish"] and len(draft["technologies"][1]) == 60 and "email" not in draft
    assert "years_of_experience" not in normalize({"years_of_experience": 400})

    extract = make_extractor(FakeClaude({"current_role": "Contable", "languages": ["Spanish"]}))
    assert extract("texto " * 80, None, None, None)["current_role"] == "Contable"
    with pytest.raises(Exception):
        make_extractor(FakeClaude({"foo": "bar"}))("texto", None, None, None)


def test_a_signed_up_person_votes_and_writes_opinions_as_themselves_and_cannot_use_owner_only_buttons(world) -> None:
    from uuid import uuid4

    world.send(OWNER_CHAT, "/invite")
    code = last(world, OWNER_CHAT)[1].split(": ")[1][:8]
    world.send("500", f"/start {code}")
    world.send("500", data="ob:consent:yes")
    with world.factory() as session:
        person = session.scalar(select(User).where(User.is_owner.is_(False)))
        person.status = "ACTIVE"
        person_id = person.id
        session.commit()
    job_id = uuid4()

    assert world.send("500", data=f"up:{job_id}") == "feedback"
    assert world.votes == [(person_id, job_id, "SAVED")]  # recorded while acting as that person, never as the owner
    assert world.send("500", "no me gusta el sueldo", reply_to=900) == "note"
    assert world.notes == [(person_id, 900, "no me gusta el sueldo")]
    before = len(world.votes)
    assert world.send("500", data=f"gl:{job_id}") == "signup_ignored" and len(world.votes) == before


def test_letter_and_interview_buttons_write_for_the_person_and_a_quota_message_is_shown(world) -> None:
    from uuid import uuid4

    from ai_job_hunter.services.person_documents import PersonDocument, QuotaExceeded

    class Docs:
        def __init__(self):
            self.calls, self.refuse = [], False

        def letter(self, user_id, job_id, language="auto"):
            self.calls.append(("letter", user_id, job_id, language))
            if self.refuse:
                raise QuotaExceeded("Ya has usado tus 3 cartas de hoy. Vuelve a probar más adelante.")
            return PersonDocument("LETTER", "Acme", "Contable", "es", "Hola equipo. Un saludo,", reused=False)

        def interview(self, user_id, job_id):
            self.calls.append(("interview", user_id, job_id))
            raise RuntimeError("provider detail that must not leak")

    docs = Docs()
    world.handlers._documents = docs
    world.send(OWNER_CHAT, "/invite")
    code = last(world, OWNER_CHAT)[1].split(": ")[1][:8]
    world.send("500", f"/start {code}")
    world.send("500", data="ob:consent:yes")
    with world.factory() as session:
        person = session.scalar(select(User).where(User.is_owner.is_(False)))
        person.status = "ACTIVE"
        person_id = person.id
        session.commit()
    job_id = uuid4()

    assert world.send("500", data=f"cl:{job_id}") == "document_written"
    assert docs.calls == [("letter", person_id, job_id, "auto")] and "Hola equipo" in last(world, 500)[1]
    assert "añade tu nombre" in last(world, 500)[1]
    docs.refuse = True
    assert world.send("500", data=f"cl:{job_id}") == "document_refused" and "3 cartas de hoy" in last(world, 500)[1]
    assert world.send("500", data=f"ip:{job_id}") == "document_failed"
    assert "provider detail" not in last(world, 500)[1]
